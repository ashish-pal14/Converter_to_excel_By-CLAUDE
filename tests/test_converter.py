"""Run with:  pytest -q

The sample PDFs are generated on first run (tests/make_samples.py), so the
suite needs no binary fixtures in the repository. Scanned tests compare OCR
output with the digital original of the same page, cell by cell.
"""
from __future__ import annotations

import datetime as dt
import io
import os
import time

import pytest
from openpyxl import load_workbook

from app.converter import ConversionError, ConvertRequest, Options, convert
from app.converter.cleaning import convert_value, parse_date, parse_number

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLES = os.path.join(ROOT, "samples")


@pytest.fixture(scope="session", autouse=True)
def samples():
    if not os.path.exists(os.path.join(SAMPLES, "password_protected.pdf")):
        from tests import make_samples
        make_samples.main()
    return SAMPLES


def run(name, **kw):
    opts = kw.pop("options", Options(layout="tables"))
    return convert(ConvertRequest(path=os.path.join(SAMPLES, name), filename=name, options=opts, **kw))


def grid(result, sheet=None):
    """Rows of a preview sheet; by default the first extracted table (not Transactions)."""
    sheets = result.preview
    if sheet is None:
        pick = next((s for s in sheets if s["name"] != "Transactions"), sheets[0])
    elif isinstance(sheet, str):
        pick = next(s for s in sheets if s["name"] == sheet)
    else:
        pick = sheets[sheet]
    return [[c["v"] for c in row] for row in pick["rows"]]


def accuracy(truth, got):
    n = ok = 0
    for tr, gr in zip(truth, got):
        for j in range(max(len(tr), len(gr))):
            a = tr[j] if j < len(tr) else ""
            b = gr[j] if j < len(gr) else ""
            if a == "" and b == "":
                continue
            n += 1
            ok += a == b
    return ok / max(n, 1)


# --------------------------------------------------------------------------- cleaning
@pytest.mark.parametrize("text,value", [
    ("1,234.56", 1234.56), ("(1,234.56)", -1234.56), ("-42", -42), ("42-", -42),
    ("$ 1,000", 1000), ("1.234,56", 1234.56), ("12.5%", 0.125), ("€99.90", 99.9),
    ("150.00 DR", -150.0), ("150.00 CR", 150.0), ("1 234 567", 1234567),
])
def test_numbers(text, value):
    assert parse_number(text)[0] == pytest.approx(value)


@pytest.mark.parametrize("text", ["00123", "000123456789012345", "Invoice 12", "abc", ""])
def test_not_numbers(text):
    assert parse_number(text) is None


def test_ocr_digit_fixes():
    assert parse_number("1O5.00", ocr=True)[0] == 105.0
    assert parse_number("21.4%.", ocr=True)[0] == pytest.approx(0.214)
    assert parse_number("1O5.00", ocr=False) is None


def test_dates():
    assert parse_date("03/04/2024", "DMY") == dt.date(2024, 4, 3)
    assert parse_date("03/04/2024", "MDY") == dt.date(2024, 3, 4)
    assert parse_date("25/12/2024", "MDY") == dt.date(2024, 12, 25)   # unambiguous
    assert parse_date("2024-01-31") == dt.date(2024, 1, 31)
    assert parse_date("5 Jan 2024") == dt.date(2024, 1, 5)
    assert parse_date("March 7, 2023") == dt.date(2023, 3, 7)
    assert parse_date("31/02/2024") is None


def test_plain_integers_have_no_separators():
    assert convert_value("2024") == (2024, "0")


# --------------------------------------------------------------------------- digital
def test_ruled_multipage_table_is_stitched_with_merged_header():
    r = run("digital_ruled_multipage.pdf")
    assert r.summary["tables"] == 1
    assert r.preview[0]["total_rows"] == 72          # 2 header rows + 70 data rows
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb.worksheets[0]
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert "C1:D1" in merged                          # "Sales" over Q1 and Q2
    assert ws.freeze_panes == "A3"
    assert isinstance(ws["C3"].value, float)          # numbers are numbers
    assert ws["E3"].number_format.endswith("%")


def test_borderless_statement():
    r = run("digital_borderless_statement.pdf")
    rows = grid(r)  # the extracted table itself
    assert rows[0] == ["Date", "Description", "Debit", "Credit", "Balance"]
    assert len(rows) == 26
    assert any("ELECTRICITY" in row[1] and "monthly" in row[1] for row in rows)  # wrapped line joined
    wb = load_workbook(io.BytesIO(r.data))
    assert isinstance(wb.worksheets[0]["A2"].value, dt.datetime)


def test_datetime_parsing():
    from app.converter.cleaning import parse_datetime
    assert parse_datetime("06 Jul 26 03:04 pm") == dt.datetime(2026, 7, 6, 15, 4)
    assert parse_datetime("06 Jul 26 03 04 pm") == dt.datetime(2026, 7, 6, 15, 4)   # OCR lost ':'
    assert parse_datetime("10 Jul 26 12:02 am") == dt.datetime(2026, 7, 10, 0, 2)
    assert parse_datetime("2026-07-06 15:04:10") == dt.datetime(2026, 7, 6, 15, 4, 10)
    assert parse_datetime("06 Jul 26 12") is None


def test_text_only_document_still_produces_a_sheet():
    r = run("digital_text_only.pdf")
    assert r.summary["tables"] == 0
    assert r.preview and r.preview[0]["rows"]


def test_password():
    with pytest.raises(ConversionError, match="password"):
        run("password_protected.pdf")
    with pytest.raises(ConversionError, match="incorrect"):
        run("password_protected.pdf", password="nope")
    assert run("password_protected.pdf", password="secret").summary["tables"] == 1


def test_page_selection():
    r = run("digital_ruled_multipage.pdf", pages=[2])
    assert r.summary["pages"] == 1


# --------------------------------------------------------------------------- statements
@pytest.mark.parametrize("name", ["digital_borderless_statement.pdf", "scanned_borderless_statement.pdf"])
def test_transactions_sheet_with_balance_check(name):
    r = run(name)
    tx = r.summary["transactions"]
    assert tx["count"] == 25
    assert tx["checks"] == ["Running balance matches on 24 of 24 rows"]
    rows = grid(r, "Transactions")
    assert rows[0] == ["Date", "Description", "Debit", "Credit", "Balance", "Page"]
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb["Transactions"]
    assert any(str(c.value).startswith("=SUM(") for row in ws.iter_rows() for c in row)


def test_card_statement_drawn_as_shapes():
    """No text layer, white-on-dark header, credits only in teal, time under date,
    repeated header/footer, three pages: like 'Print to PDF' card statements."""
    import json
    truth = json.load(open(os.path.join(SAMPLES, "card_statement_vector_text.json")))
    r = run("card_statement_vector_text.pdf")
    assert all(p["type"] == "scanned" for p in r.summary["page_details"])
    tx = r.summary["transactions"]
    assert tx["count"] == len(truth)
    assert "text colour" in tx["method"]
    exp_debits = round(sum(a for _, _, a, credit in truth if not credit), 2)
    exp_credits = round(sum(a for _, _, a, credit in truth if credit), 2)
    assert tx["debits"] == pytest.approx(exp_debits, abs=0.01)
    assert tx["credits"] == pytest.approx(exp_credits, abs=0.01)
    rows = grid(r, "Transactions")
    assert rows[0][:3] == ["Date", "Time", "Description"]
    assert [row[0] for row in rows[1:len(truth) + 1]] == [d for d, *_ in truth]
    # repeated header/footer kept once, not on every page
    names = [s["name"] for s in r.preview]
    assert sum(1 for n in names if n.startswith("Table")) <= 4


# --------------------------------------------------------------------------- scanned
@pytest.mark.parametrize("scan,truth,pages,minimum", [
    ("scanned_ruled_skewed.pdf", "digital_ruled_multipage.pdf", [1], 0.97),
    ("photo_of_table.png", "digital_ruled_multipage.pdf", [1], 0.97),
    ("scanned_borderless_statement.pdf", "digital_borderless_statement.pdf", None, 0.97),
    ("scanned_rotated_90.pdf", "digital_borderless_statement.pdf", None, 0.97),
])
def test_scanned_accuracy(scan, truth, pages, minimum):
    expected = grid(run(truth, pages=pages))
    got = run(scan)
    assert got.summary["scanned_pages"] == 1
    assert accuracy(expected, grid(got)) >= minimum


def test_mixed_page_ocrs_embedded_image():
    r = run("mixed_digital_with_scan.pdf")
    assert r.summary["digital_pages"] == 1
    assert r.summary["tables"] >= 1
    assert "OCR used on an embedded image" in r.summary["page_details"][0]["notes"]


def test_output_formats():
    assert run("digital_borderless_statement.pdf", output="csv").extension == "zip"
    assert run("digital_borderless_statement.pdf", output="json").data.startswith(b"{")


# --------------------------------------------------------------------------- exact copy
def test_exact_copy_is_the_default_and_keeps_layout():
    assert Options().layout == "exact"
    r = run("digital_borderless_statement.pdf", options=Options())
    assert r.summary["mode"] == "exact"
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Page 1"]
    ws = wb["Page 1"]
    cells = {c.value: c for row in ws.iter_rows() for c in row if c.value is not None}
    assert "Acme Bank - Account Statement" in cells
    title = cells["Acme Bank - Account Statement"]
    assert title.font.bold and title.font.size >= 13
    # amounts are numbers but look exactly as printed
    bal = cells[5768.84]
    assert bal.number_format == "#,##0.00" and bal.alignment.horizontal == "right"
    # dates are real dates in the printed format
    first = next(c for c in cells.values() if isinstance(c.value, dt.datetime))
    assert first.number_format == "dd/mm/yyyy"
    # the account number is not turned into a (rounded) number
    assert any(isinstance(v, str) and "000123456789012345" in v for v in cells)
    # rows follow the page from top to bottom
    rows = [c.row for c in cells.values()]
    assert title.row == min(rows)


def test_exact_copy_values_can_stay_text():
    r = run("digital_borderless_statement.pdf", options=Options(convert_types=False))
    ws = load_workbook(io.BytesIO(r.data)).active
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "5,768.84" in values and all(not isinstance(v, (int, float)) for v in values)


def test_exact_copy_all_pages_on_one_sheet():
    r = run("digital_ruled_multipage.pdf", options=Options(layout="exact_single"))
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["All pages"]
    ws = wb.active
    assert len(ws.row_breaks.brk) == 1                 # page break between the 2 pages
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert merged                                        # ruled table merged header kept


def test_exact_copy_of_card_statement_drawn_as_shapes():
    import json
    truth = json.load(open(os.path.join(SAMPLES, "card_statement_vector_text.json")))
    r = run("card_statement_vector_text.pdf", options=Options())
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Page 1", "Page 2", "Page 3"]
    amounts, teal = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, (int, float)) and c.number_format.startswith('"Rs'):
                    amounts.append(round(c.value, 2))
                    color = c.font.color.rgb if c.font.color is not None else None
                    if isinstance(color, str) and color[-6:-4] < "80" and color[-4:-2] > "80":
                        teal += 1                         # credits keep their teal colour
    assert sorted(amounts) == sorted(round(a, 2) for *_, a, _ in truth)
    assert teal == sum(1 for *_, credit in truth if credit)
    # the dark header bar is painted across the row
    ws = wb["Page 1"]
    dark = [c for row in ws.iter_rows() for c in row
            if c.fill and c.fill.fgColor and isinstance(c.fill.fgColor.rgb, str)
            and c.fill.fgColor.rgb[-6:] in ("3F3F3F", "404040", "3E3E3E", "414141", "3D3D3D")]
    assert len(dark) >= 3


def test_exact_copy_preview_and_csv():
    r = run("digital_borderless_statement.pdf", options=Options(), output="csv")
    assert r.extension == "zip"
    assert r.preview[0]["name"] == "Page 1" and r.preview[0]["rows"]


# --------------------------------------------------------------------------- API
def test_api_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    from fastapi.testclient import TestClient
    import importlib
    import app.config
    import app.main
    importlib.reload(app.config)
    importlib.reload(app.main)
    with TestClient(app.main.app) as client:
        assert client.get("/api/health").json() == {"status": "ok"}
        with open(os.path.join(SAMPLES, "digital_borderless_statement.pdf"), "rb") as f:
            job = client.post("/api/jobs", files={"file": ("s.pdf", f, "application/pdf")}).json()
        for _ in range(100):
            status = client.get(f"/api/jobs/{job['id']}").json()
            if status["status"] in ("done", "failed"):
                break
            time.sleep(0.1)
        assert status["status"] == "done"
        assert client.get(f"/api/jobs/{job['id']}/download").status_code == 200
        assert client.get(f"/api/jobs/{job['id']}/preview").json()["sheets"]

        bad = client.post("/api/jobs", files={"file": ("x.txt", b"hello", "text/plain")})
        assert bad.status_code == 415
        bad = client.post("/api/jobs", files={"file": ("s.pdf", b"%PDF-1.4", "application/pdf")},
                          data={"pages": "x"})
        assert bad.status_code == 422
