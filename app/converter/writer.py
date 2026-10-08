"""Write converted pages to .xlsx, .csv (zipped) or .json."""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .cleaning import convert_value, looks_like_header
from .models import PageResult, Table, TextBlock
from .statement import build_transactions

HEADER_FILL = PatternFill("solid", start_color="E3ECE6")
LOW_CONF_FILL = PatternFill("solid", start_color="FFF2B3")
TITLE_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True)
THIN = Side(style="thin", color="C9D3CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass
class Options:
    layout: str = "exact"           # exact | exact_single | tables | pages | single
    date_order: str = "DMY"         # DMY | MDY
    decimal: str = "auto"           # auto | . | ,
    convert_types: bool = True
    include_text: bool = True
    low_conf_threshold: float = 60.0
    statement: str = "auto"         # auto: add a Transactions sheet when statement tables are found; off


@dataclass
class SheetData:
    name: str
    rows: List[List[Any]] = field(default_factory=list)          # typed values
    formats: List[List[Optional[str]]] = field(default_factory=list)
    flags: List[List[int]] = field(default_factory=list)          # 0 ok, 1 low conf, 2 header, 3 title
    merges: List[Tuple[int, int, int, int]] = field(default_factory=list)
    comments: Dict[Tuple[int, int], str] = field(default_factory=dict)
    formulas: Dict[Tuple[int, int], str] = field(default_factory=dict)  # written instead of the value
    colors: Dict[Tuple[int, int], str] = field(default_factory=dict)    # original ink colour
    freeze_header: int = 0          # number of header rows to freeze


# --------------------------------------------------------------------------- building
def _safe_name(name: str, used: set) -> str:
    name = re.sub(r"[\[\]\:\*\?\/\\]", "-", name).strip("'") or "Sheet"
    name = name[:31]
    base, k = name, 2
    while name.lower() in used:
        suffix = f" ({k})"
        name = base[: 31 - len(suffix)] + suffix
        k += 1
    used.add(name.lower())
    return name


def _table_label(t: Table) -> str:
    pages = t.pages or [t.page]
    return f"p{pages[0]}" if len(pages) == 1 else f"p{pages[0]}-{pages[-1]}"


def _add_table(sheet: SheetData, t: Table, opts: Options, start_row: int,
               title: Optional[str] = None) -> int:
    r = start_row
    if title:
        _put_row(sheet, r, [title], [None], [3])
        r += 1
    is_ocr = t.source.startswith("ocr")
    header_rows = 0
    while (header_rows < min(3, len(t.rows) - 1)
           and looks_like_header(t.rows[header_rows])):
        header_rows += 1
    if header_rows == len(t.rows):
        header_rows = 0
    for i, row in enumerate(t.rows):
        values, fmts, flags = [], [], []
        is_header = i < header_rows
        for j, cell in enumerate(row):
            if opts.convert_types and not is_header:
                v, f = convert_value(cell, opts.date_order, opts.decimal, ocr=is_ocr)
            else:
                v, f = (cell or "").strip(), None
            flag = 2 if is_header else 0
            conf = None
            if t.confidence is not None and i < len(t.confidence) and j < len(t.confidence[i]):
                conf = t.confidence[i][j]
            if t.colors is not None and i < len(t.colors) and j < len(t.colors[i]) and t.colors[i][j]:
                sheet.colors[(r, j)] = t.colors[i][j]
            if conf is not None and v != "" and conf < opts.low_conf_threshold:
                flag = 1
                sheet.comments[(r, j)] = f"Low OCR confidence ({conf:.0f}%). Please check this value."
            values.append(v)
            fmts.append(f)
            flags.append(flag)
        _put_row(sheet, r, values, fmts, flags)
        r += 1
    for r0, c0, r1, c1 in t.merges:
        sheet.merges.append((start_row + (1 if title else 0) + r0, c0,
                             start_row + (1 if title else 0) + r1, c1))
    return r


def _add_text(sheet: SheetData, block: TextBlock, start_row: int) -> int:
    r = start_row
    for line in block.lines:
        _put_row(sheet, r, [line], [None], [0])
        r += 1
    return r


def _put_row(sheet: SheetData, r: int, values, fmts, flags):
    while len(sheet.rows) <= r:
        sheet.rows.append([])
        sheet.formats.append([])
        sheet.flags.append([])
    sheet.rows[r] = list(values)
    sheet.formats[r] = list(fmts)
    sheet.flags[r] = list(flags)


def _transactions_sheet(pages: List[PageResult], opts: Options, used: set):
    tx = build_transactions(pages, opts.date_order)
    if tx is None:
        return None, None
    s = SheetData(_safe_name("Transactions", used))
    _put_row(s, 0, tx.header, [None] * len(tx.header), [2] * len(tx.header))
    for i, (vals, fmts, flags) in enumerate(zip(tx.rows, tx.formats, tx.flags), start=1):
        _put_row(s, i, vals, fmts, flags)
    for (r, c), text in tx.comments.items():
        s.comments[(r + 1, c)] = text
    for r in range(1, len(tx.rows) + 1):
        for c in (tx.debit_col, tx.credit_col):
            if s.flags[r][c] == 1 and (r, c) not in s.comments:
                s.comments[(r, c)] = "Low OCR confidence. Please check this amount."
    n = len(tx.rows)
    last = n  # 0-based index of last data row
    total_row = n + 2
    desc_col = tx.header.index("Description")
    dl = get_column_letter(tx.debit_col + 1)
    cl = get_column_letter(tx.credit_col + 1)
    def num(v):
        return v if isinstance(v, (int, float)) else 0
    debits = sum(num(r[tx.debit_col]) for r in tx.rows)
    credits = sum(num(r[tx.credit_col]) for r in tx.rows)
    totals = [
        ("Total debits", tx.debit_col, f"=SUM({dl}2:{dl}{last + 1})", debits),
        ("Total credits", tx.credit_col, f"=SUM({cl}2:{cl}{last + 1})", credits),
    ]
    for k, (label, col, formula, value) in enumerate(totals):
        r = total_row + k
        vals = [""] * len(tx.header)
        fl = [0] * len(tx.header)
        fm: List[Optional[str]] = [None] * len(tx.header)
        vals[desc_col], fl[desc_col] = label, 3
        vals[col], fm[col], fl[col] = round(value, 2), "#,##0.00", 3
        _put_row(s, r, vals, fm, fl)
        s.formulas[(r, col)] = formula
    r = total_row + 2
    vals = [""] * len(tx.header)
    vals[desc_col] = "Transactions"
    vals[tx.debit_col] = n
    _put_row(s, r, vals, [None] * len(tx.header), [0] * len(tx.header))
    s.formulas[(r, tx.debit_col)] = f"=COUNTA(A2:A{last + 1})"
    s.freeze_header = 1
    info = {"count": n, "debits": round(debits, 2), "credits": round(credits, 2),
            "method": tx.method, "checks": tx.checks}
    return s, info


def build_sheets(pages: List[PageResult], opts: Options, info: Optional[dict] = None) -> List[SheetData]:
    used: set = set()
    sheets: List[SheetData] = []
    if opts.statement != "off":
        tx_sheet, tx_info = _transactions_sheet(pages, opts, used)
        if tx_sheet is not None:
            sheets.append(tx_sheet)
            if info is not None:
                info["transactions"] = tx_info
    all_tables = [t for p in pages for t in p.tables]
    layout = opts.layout
    if layout == "tables" and not all_tables:
        layout = "pages"   # nothing tabular: fall back so the file is never empty

    if layout == "tables":
        for n, t in enumerate(all_tables, 1):
            s = SheetData(_safe_name(f"Table {n} ({_table_label(t)})", used))
            _add_table(s, t, opts, 0)
            s.freeze_header = sum(1 for f in s.flags if f and all(x == 2 for x in f))
            sheets.append(s)
        if opts.include_text:
            s = SheetData(_safe_name("Text", used))
            r = 0
            for p in pages:
                texts = [b for b in p.blocks if isinstance(b, TextBlock)]
                if not texts:
                    continue
                _put_row(s, r, [f"Page {p.number}"], [None], [3])
                r += 1
                for b in texts:
                    r = _add_text(s, b, r) + 1
            if s.rows:
                sheets.append(s)

    elif layout == "pages":
        for p in pages:
            s = SheetData(_safe_name(f"Page {p.number}", used))
            r = 0
            for b in p.blocks:
                if isinstance(b, Table):
                    r = _add_table(s, b, opts, r) + 1
                elif opts.include_text:
                    r = _add_text(s, b, r) + 1
            if not s.rows:
                _put_row(s, 0, ["(no text found on this page)"], [None], [0])
            sheets.append(s)

    else:  # single
        s = SheetData(_safe_name("All tables", used))
        r = 0
        for n, t in enumerate(all_tables, 1):
            r = _add_table(s, t, opts, r, title=f"Table {n} (page {_table_label(t)[1:]})") + 1
        if not all_tables:
            for p in pages:
                for b in p.blocks:
                    r = _add_text(s, b, r)
        sheets.append(s)
    return sheets


def _summary_rows(pages: List[PageResult], meta: dict) -> List[List[Any]]:
    rows: List[List[Any]] = [
        ["Source file", meta.get("filename", "")],
        ["Converted", dt.datetime.now().strftime("%Y-%m-%d %H:%M")],
        ["Pages", len(pages)],
        ["Tables found", sum(len(p.tables) for p in pages)],
    ]
    tx = meta.get("transactions")
    if tx:
        rows += [["Transactions", tx["count"]],
                 ["Total debits", tx["debits"]],
                 ["Total credits", tx["credits"]],
                 ["Debit/credit from", tx["method"]]]
        rows += [["Check", c] for c in tx["checks"]]
    rows += [
        [],
        ["Page", "Type", "Tables", "OCR confidence", "Notes"],
    ]
    for p in pages:
        rows.append([p.number, p.kind, len(p.tables),
                     "" if p.mean_confidence is None else f"{p.mean_confidence:.0f}%",
                     "; ".join(p.notes)])
    rows += [[], ["Yellow cells were read by OCR with low confidence. Check them against the PDF."]]
    return rows


# --------------------------------------------------------------------------- writers
def write_xlsx(sheets: List[SheetData], pages: List[PageResult], meta: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for s in sheets:
        ws = wb.create_sheet(s.name)
        widths: Dict[int, float] = {}
        for r, row in enumerate(s.rows):
            for c, v in enumerate(row):
                if v == "" or v is None:
                    continue
                cell = ws.cell(row=r + 1, column=c + 1, value=s.formulas.get((r, c), v))
                ink = s.colors.get((r, c))
                fmt = s.formats[r][c] if c < len(s.formats[r]) else None
                flag = s.flags[r][c] if c < len(s.flags[r]) else 0
                if fmt:
                    cell.number_format = fmt
                if flag == 1:
                    cell.fill = LOW_CONF_FILL
                elif flag == 2:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                    cell.border = BORDER
                elif flag == 3:
                    cell.font = TITLE_FONT
                if ink and flag != 2:
                    cell.font = Font(color=ink, bold=flag == 3)
                if isinstance(v, str) and len(v) > 60:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                comment = s.comments.get((r, c))
                if comment:
                    cell.comment = Comment(comment, "PDF to Excel")
                length = len(v.strftime("%Y-%m-%d")) if isinstance(v, (dt.date, dt.datetime)) else len(str(v))
                if flag != 3:
                    widths[c] = max(widths.get(c, 0), min(length, 60))
        taken = set()
        for r0, c0, r1, c1 in s.merges:
            cells = {(r, c) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)}
            if cells & taken or (r0 == r1 and c0 == c1):
                continue
            taken |= cells
            ws.merge_cells(start_row=r0 + 1, start_column=c0 + 1, end_row=r1 + 1, end_column=c1 + 1)
            ws.cell(row=r0 + 1, column=c0 + 1).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True)
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c + 1)].width = max(8, w + 2)
        if s.freeze_header:
            ws.freeze_panes = f"A{int(s.freeze_header) + 1}"

    ws = wb.create_sheet("About this file")
    for r, row in enumerate(_summary_rows(pages, meta), 1):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    for row in ws.iter_rows():
        if row[0].value == "Page" and len(row) > 1 and row[1].value == "Type":
            for c in row:
                c.font = HEADER_FONT
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _plain(v: Any) -> Any:
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()
    return v


def write_csv_zip(sheets: List[SheetData]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for s in sheets:
            out = io.StringIO()
            w = csv.writer(out)
            for row in s.rows:
                w.writerow([_plain(v) for v in row])
            zf.writestr(re.sub(r"[^\w\- ()]", "_", s.name) + ".csv", "\ufeff" + out.getvalue())
    return buf.getvalue()


def write_json(sheets: List[SheetData], pages: List[PageResult], meta: dict) -> bytes:
    doc = {
        "source": meta.get("filename"),
        "pages": [{"page": p.number, "type": p.kind, "notes": p.notes,
                   "ocr_confidence": p.mean_confidence} for p in pages],
        "sheets": [{"name": s.name, "rows": [[_plain(v) for v in row] for row in s.rows]}
                   for s in sheets],
    }
    return json.dumps(doc, ensure_ascii=False, indent=2).encode("utf-8")


def _display(v: Any, fmt: Optional[str]) -> str:
    """Render a value roughly the way Excel will show it."""
    if v is None or v == "":
        return ""
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) or isinstance(v, int):
        if fmt and fmt.endswith("%"):
            decimals = fmt.count("0") - 1
            return f"{v * 100:.{max(decimals, 0)}f}%"
        if fmt and "#,##0" in fmt:
            decimals = len(fmt.split(".")[1].split(";")[0].rstrip('"')) if "." in fmt else 0
            prefix = fmt.split('"')[1] if fmt.startswith('"') else ""
            s = f"{abs(v):,.{decimals}f}"
            return ("-" if v < 0 else "") + prefix + s
        return f"{v:.10g}"
    return str(v)


def preview(sheets: List[SheetData], max_rows: int = 60, max_cols: int = 20) -> list:
    out = []
    for s in sheets:
        rows = []
        for r, row in enumerate(s.rows[:max_rows]):
            rows.append([
                {"v": _display(v, s.formats[r][c] if c < len(s.formats[r]) else None),
                 "f": s.flags[r][c] if c < len(s.flags[r]) else 0,
                 "n": isinstance(v, (int, float)) and not isinstance(v, bool)}
                for c, v in enumerate(row[:max_cols])
            ])
        out.append({"name": s.name, "rows": rows, "total_rows": len(s.rows)})
    return out
