"""Exact copy: lay each page out in Excel the way it looks in the PDF.

Every line of the page becomes a row and every piece of text goes into the
column that matches its horizontal position. Row heights and column widths
are taken from the page's own measurements, so the sheet reads like the PDF:

  * rows       one per visual line; height = distance to the next line
  * columns    boundaries where text starts (left-aligned text), where
               amounts end (right-aligned numbers) and where ruled-table
               cells begin and end
  * formatting font size, bold, text colour, background bands, table borders
               and merged cells
  * values     numbers and dates become real Excel values but are formatted
               to look exactly as printed ("Rs. 1,000", "06 Jul 26")
"""
from __future__ import annotations

import datetime as dt
import io
import re
import statistics
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.worksheet.pagebreak import Break

from .cleaning import parse_date, parse_datetime, parse_number
from .models import PageResult, Segment

LOW_CONF = 60.0
THIN = Side(style="thin", color="7F8A84")
CHECK_FILL = "FFF2B3"
FONT_NAME = "Arial"           # metrics close to Helvetica, the usual statement font


# --------------------------------------------------------------------------- values
_DATE_FORMATS = [
    (re.compile(r"^\d{1,2} [A-Za-z]{3} \d{2}$"), "dd mmm yy"),
    (re.compile(r"^\d{1,2} [A-Za-z]{3} \d{4}$"), "dd mmm yyyy"),
    (re.compile(r"^\d{1,2} [A-Za-z]{4,} \d{4}$"), "dd mmmm yyyy"),
    (re.compile(r"^\d{1,2}-[A-Za-z]{3}-\d{2}$"), "dd-mmm-yy"),
    (re.compile(r"^\d{1,2}-[A-Za-z]{3}-\d{4}$"), "dd-mmm-yyyy"),
    (re.compile(r"^[A-Za-z]{3} \d{1,2}, \d{4}$"), "mmm dd, yyyy"),
    (re.compile(r"^[A-Za-z]{4,} \d{1,2}, \d{4}$"), "mmmm dd, yyyy"),
    (re.compile(r"^\d{4}-\d{2}-\d{2}$"), "yyyy-mm-dd"),
]


def _numeric_date_format(t: str, order: str) -> Optional[str]:
    m = re.match(r"^(\d{1,2})([/.\-])(\d{1,2})\2(\d{2}|\d{4})$", t)
    if not m:
        return None
    sep = m.group(2)
    year = "yyyy" if len(m.group(4)) == 4 else "yy"
    d = "dd" if len(m.group(1 if order == "DMY" else 3)) == 2 else "d"
    mo = "mm" if len(m.group(3 if order == "DMY" else 1)) == 2 else "m"
    first, second = (d, mo) if order == "DMY" else (mo, d)
    return f"{first}{sep}{second}{sep}{year}"


def _date_format(date_text: str, order: str) -> Optional[str]:
    for rx, fmt in _DATE_FORMATS:
        if rx.match(date_text):
            return fmt
    return _numeric_date_format(date_text, order)


def _quote(s: str) -> str:
    return '"' + s.replace('"', "") + '"' if s else ""


def _number_format(text: str, value: float) -> Optional[str]:
    """Excel format that shows value exactly like text (prefix, separators, decimals)."""
    t = text.strip()
    first = re.search(r"\d", t)
    last = max(i for i, ch in enumerate(t) if ch.isdigit())
    prefix, core, suffix = t[:first.start()], t[first.start():last + 1], t[last + 1:]
    paren = prefix.strip().startswith("(") and suffix.strip().endswith(")")
    prefix = prefix.replace("(", "").replace("-", "").replace("−", "").replace("+", "")
    suffix = suffix.replace(")", "")
    pct = suffix.strip().startswith("%")
    if pct:
        suffix = suffix.replace("%", "", 1)
    trailing_minus = suffix.strip() == "-"
    if trailing_minus:
        suffix = ""
    decimals = 0
    m = re.search(r"[.,](\d+)$", core)
    if m and not re.fullmatch(r"\d{1,3}(,\d{3})+", core):
        decimals = len(m.group(1))
    dec = ("." + "0" * decimals) if decimals else ""
    if re.search(r"\d,\d{2},\d{3}", core):          # Indian grouping 1,00,000
        base = (f"[>=10000000]##\\,##\\,##\\,##0{dec};[>=100000]##\\,##\\,##0{dec};##,##0{dec}")
        if value < 0 or prefix.strip() or suffix.strip() or pct:
            base = f"#,##0{dec}"
        else:
            return base
    elif "," in core and "." in core and core.rfind(",") > core.rfind("."):
        return None                                   # 1.234,56: locale-specific, keep as text
    else:
        base = ("#,##0" if "," in core else "0") + dec
    if pct:
        base += "%"
    pos = f"{_quote(prefix)}{base}{_quote(suffix)}"
    if paren:
        return f"{pos};({pos})"
    if trailing_minus:
        return f"{pos};{pos}-"
    if suffix.strip().upper() in ("DR", "DR."):
        return f"{pos};{pos}"                        # value is negative, shown as "... DR"
    if prefix.strip():
        return f"{pos};-{pos}"
    return pos


def display_value(text: str, order: str = "DMY", ocr: bool = False):
    """(value, number_format) that displays like text, or (text, None) to keep text."""
    t = text.strip()
    if not t or len(t) > 40:
        return text, None
    d = parse_date(t, order)
    if d is not None:
        fmt = _date_format(t, order)
        return (d, fmt) if fmt else (text, None)
    dtm = parse_datetime(t, order)
    if dtm is not None:
        m = re.match(r"^(.*?)[\s,T]+\d{1,2}\s*[:.]?\s*\d{2}", t)
        fmt = _date_format(m.group(1).strip(" ,:"), order) if m else None
        if fmt:
            ampm = re.search(r"[AaPp]\.?\s?[Mm]\.?$", t)
            return dtm, fmt + (" hh:mm AM/PM" if ampm else " hh:mm")
        return text, None
    p = parse_number(t, ocr=ocr)
    if p is None:
        return text, None
    value = p[0]
    if not any(ch.isdigit() for ch in t):
        return text, None
    fmt = _number_format(t, value)
    if fmt is None:
        return text, None
    if float(value).is_integer() and "." not in fmt.split(";")[0]:
        value = int(value)
    return value, fmt


# --------------------------------------------------------------------------- layout
@dataclass
class Placed:
    seg: Segment
    row: int
    col: int
    row_end: int
    col_end: int
    align: str


@dataclass
class ExactSheet:
    name: str
    col_widths_pt: List[float] = field(default_factory=list)
    row_heights_pt: List[float] = field(default_factory=list)
    cells: List[Placed] = field(default_factory=list)
    page_break_rows: List[int] = field(default_factory=list)
    # background areas as cell ranges: (row0, col0, row1, col1, colour), 0-based like cells
    fills: List[Tuple[int, int, int, int, str]] = field(default_factory=list)
    # horizontal rules as bottom borders: (row, col0, col1, colour)
    rules: List[Tuple[int, int, int, str]] = field(default_factory=list)
    # (row, row offset pt, col, col offset pt, width pt, height pt, png)
    pictures: List[Tuple[int, float, int, float, float, float, bytes]] = field(default_factory=list)
    ocr: bool = False


def _cluster(values: List[float], tol: float) -> List[float]:
    values = sorted(values)
    groups: List[List[float]] = []
    for v in values:
        if groups and v - groups[-1][-1] <= tol:
            groups[-1].append(v)
        else:
            groups.append([v])
    return [statistics.median(g) for g in groups]


def _nearest(bounds: List[float], v: float) -> int:
    return min(range(len(bounds)), key=lambda i: abs(bounds[i] - v))


def _right_aligned(seg: Segment) -> bool:
    """Amounts are printed right-aligned; phone numbers, IDs and dates are not."""
    if seg.align:
        return seg.align == "right"
    t = seg.text.strip()
    if not t or parse_date(t) or parse_datetime(t):
        return False
    p = parse_number(t)
    if p is None:
        return False
    digits = sum(ch.isdigit() for ch in t)
    styled = any(ch in t for ch in ".,%") or bool(re.search(r"[₹$€£¥]|\bRs\b|\bINR\b|CR$|DR$", t, re.I))
    return styled or digits <= 7


def _text_width(text: str, size: float) -> float:
    """Width Excel needs to show text without ### (Arial ~0.6 em/char + cell padding)."""
    return len(text) * size * 0.6 + 6.0


def layout_pages(pages: List[PageResult], name: str) -> ExactSheet:
    """Place the segments of one or more pages (stacked) onto one sheet."""
    sheet = ExactSheet(name)
    segs: List[Segment] = []
    fill_rects, rule_lines, pics = [], [], []
    page_tops: List[float] = []
    offset = 0.0
    width = max((p.width for p in pages), default=595.0)
    for p in pages:
        page_tops.append(offset)
        for s in p.segments:
            if s.text.strip() or s.boxed:
                segs.append(Segment(**{**s.__dict__, "top": s.top + offset, "bottom": s.bottom + offset}))
        fill_rects += [(x0, y0 + offset, x1, y1 + offset, c) for x0, y0, x1, y1, c in p.fills]
        rule_lines += [(x0, y + offset, x1, c) for x0, y, x1, c in p.rules]
        pics += [((b[0], b[1] + offset, b[2], b[3] + offset), png) for b, png in p.pictures]
        offset += p.height
    sheet.ocr = any(p.kind == "scanned" for p in pages)
    if not segs and not pics:
        sheet.col_widths_pt = [width]
        sheet.row_heights_pt = [15.0]
        return sheet

    # ---------------------------------------------------------------- rows
    free_h = [s.height for s in segs if not s.boxed] or [12.0]
    boxed_h = [s.height for s in segs if s.boxed]
    cell_h = statistics.median(boxed_h) if boxed_h else statistics.median(free_h)

    def band_of(s: Segment) -> Tuple[float, float]:
        limit = cell_h * 1.1 if s.boxed else s.height
        return s.top, s.top + min(s.height, limit)

    rows: List[List[Segment]] = []
    bands: List[Tuple[float, float]] = []
    for s in sorted(segs, key=lambda s: (band_of(s)[0], s.x0)):
        b0, b1 = band_of(s)
        for i in range(len(rows) - 1, max(-1, len(rows) - 6), -1):
            r0, r1 = bands[i]
            if min(r1, b1) - max(r0, b0) >= 0.5 * min(r1 - r0, b1 - b0):
                rows[i].append(s)
                bands[i] = (min(r0, b0), max(r1, b1))
                break
        else:
            rows.append([s])
            bands.append((b0, b1))
    order = sorted(range(len(rows)), key=lambda i: bands[i][0])
    rows = [rows[i] for i in order]
    bands = [bands[i] for i in order]
    row_tops = [min(s.top for s in r) for r in rows]
    n_rows = len(rows)

    # separator lines that sit in the gap below a row
    rule_after: Dict[int, List[Tuple[float, float, float, str]]] = {}
    for x0, y, x1, color in rule_lines:
        if any(b0 + 1 < y < b1 - 1 for b0, b1 in bands):
            continue                                    # runs through text
        above = [i for i in range(n_rows) if bands[i][1] <= y + 1]
        if above:
            rule_after.setdefault(max(above), []).append((x0, y, x1, color))

    # ---------------------------------------------------------------- excel rows
    heights: List[float] = []
    row_map: Dict[int, int] = {}          # content row -> excel row index (0-based)
    rule_rows: List[Tuple[int, float, float, str]] = []

    def page_of(y: float) -> int:
        return max(i for i, t in enumerate(page_tops) if t <= y + 0.01)

    first_top = row_tops[0] if n_rows else page_tops[0] + 20
    heights.append(max(2.0, first_top - page_tops[page_of(first_top)]))      # top margin
    for i in range(n_rows):
        size = max(s.size for s in rows[i])
        min_h = max(size * 1.2, bands[i][1] - row_tops[i] + 1)
        same_page = i + 1 < n_rows and page_of(row_tops[i + 1]) == page_of(row_tops[i])
        if i + 1 < n_rows and same_page:
            next_top = row_tops[i + 1]
        else:
            next_top = row_tops[i] + min_h + 2
        row_map[i] = len(heights)
        rules_here = rule_after.get(i, [])
        y_rule = min((r[1] for r in rules_here), default=None)
        if y_rule is not None and y_rule - row_tops[i] >= min_h * 0.9 and next_top - y_rule >= 1:
            heights.append(min(409.0, y_rule - row_tops[i]))
            for x0, _, x1, color in rules_here:
                rule_rows.append((len(heights) - 1, x0, x1, color))
            heights.append(min(409.0, max(1.0, next_top - y_rule)))           # spacer under the line
        else:
            heights.append(min(409.0, max(next_top - row_tops[i], min_h)))
            for x0, _, x1, color in rules_here:
                rule_rows.append((len(heights) - 1, x0, x1, color))
        if i + 1 < n_rows and not same_page:
            # next page: page break, then that page's top margin
            sheet.page_break_rows.append(len(heights))
            new_page = page_of(row_tops[i + 1])
            heights.append(max(2.0, row_tops[i + 1] - page_tops[new_page]))
    sheet.row_heights_pt = heights

    # ---------------------------------------------------------------- columns
    edges: List[float] = [0.0, width]
    kinds: Dict[int, str] = {}
    for x0, _, x1, _, _ in fill_rects:
        edges += [x0, x1]
    for x0, _, x1, _ in rule_lines:
        edges += [x0, x1]
    # Numbers are usually right-aligned, but not always (a "No. of days"
    # column may be left-aligned). Look at the neighbours: a number whose left
    # edge lines up with other text is left-aligned; one whose right edge lines
    # up with other numbers is right-aligned.
    numeric = [i for i, s in enumerate(segs) if not s.boxed and _right_aligned(s)]
    lefts = [s.x0 for s in segs if not s.boxed]
    rights = [segs[i].x1 for i in numeric]
    for idx, s in enumerate(segs):
        if s.boxed:
            edges += [s.x0, s.x1]
            continue
        right = idx in numeric
        if right:
            left_matches = sum(1 for x in lefts if abs(x - s.x0) <= 2.0) - 1
            right_matches = sum(1 for x in rights if abs(x - s.x1) <= 2.0) - 1
            if left_matches > right_matches:
                right = False
        if right:
            edges.append(s.x1)
            kinds[idx] = "right"
        else:
            edges.append(s.x0)
            kinds[idx] = "left"
    bounds = _cluster(edges, tol=3.0)
    if bounds[0] > 0.5:
        bounds.insert(0, 0.0)
    sheet.col_widths_pt = [max(0.5, b - a) for a, b in zip(bounds, bounds[1:])]
    n_cols = len(sheet.col_widths_pt)
    seg_index = {id(s): i for i, s in enumerate(segs)}

    def col_span(x0: float, x1: float) -> Tuple[int, int]:
        c0 = min(_nearest(bounds, x0), n_cols - 1)
        return c0, max(c0, min(_nearest(bounds, x1) - 1, n_cols - 1))

    # ---------------------------------------------------------------- backgrounds and lines
    for x0, y0, x1, y1, color in sorted(fill_rects, key=lambda f: -(f[2] - f[0]) * (f[3] - f[1])):
        c0, c1 = col_span(x0, x1)
        inside = [i for i, (b0, b1) in enumerate(bands) if min(b1, y1) - max(b0, y0) >= 0.5 * (b1 - b0)]
        if inside:
            sheet.fills.append((row_map[min(inside)], c0, row_map[max(inside)], c1, color))
    for r, x0, x1, color in rule_rows:
        c0, c1 = col_span(x0, x1)
        sheet.rules.append((r, c0, c1, color))

    # ---------------------------------------------------------------- place text
    taken: Dict[Tuple[int, int], Placed] = {}
    occupied: Dict[Tuple[int, int], Placed] = {}
    for i, r in enumerate(rows):
        er = row_map[i]
        for s in sorted(r, key=lambda s: s.x0):
            k = seg_index[id(s)]
            if s.boxed:
                c0, c1 = col_span(s.x0, s.x1)
                last = i
                for j in range(i + 1, n_rows):
                    if row_tops[j] < s.bottom - 0.4 * cell_h:
                        last = j
                    else:
                        break
                placed = Placed(s, er, c0, row_map[last], c1, s.align or ("right" if _right_aligned(s) else "left"))
            elif kinds.get(k) == "right":
                c = max(0, min(_nearest(bounds, s.x1) - 1, n_cols - 1))
                placed = Placed(s, er, c, er, c, "right")
            else:
                c = min(_nearest(bounds, s.x0), n_cols - 1)
                placed = Placed(s, er, c, er, c, "left")
            key = (placed.row, placed.col)
            if key in taken:
                other = taken[key]
                other.seg = Segment(**{**other.seg.__dict__,
                                       "text": (other.seg.text + " " + s.text).strip(),
                                       "conf": min(other.seg.conf, s.conf)})
                continue
            taken[key] = placed
            for rr in range(placed.row, placed.row_end + 1):
                for cc in range(placed.col, placed.col_end + 1):
                    occupied[(rr, cc)] = placed
            sheet.cells.append(placed)

    # numbers and dates never spill into the next cell (Excel shows ###):
    # give them the empty cells beside them
    for p in sheet.cells:
        if p.seg.boxed or p.row != p.row_end:
            continue
        value, _ = display_value(p.seg.text)
        if isinstance(value, str):
            continue
        need = _text_width(p.seg.text, p.seg.size)

        def have():
            return sum(sheet.col_widths_pt[p.col:p.col_end + 1])
        if p.align == "right":
            while have() < need and p.col > 0 and (p.row, p.col - 1) not in occupied:
                p.col -= 1
                occupied[(p.row, p.col)] = p
        else:
            while have() < need and p.col_end + 1 < n_cols and (p.row, p.col_end + 1) not in occupied:
                p.col_end += 1
                occupied[(p.row, p.col_end)] = p

    # ---------------------------------------------------------------- pictures
    ycum = [0.0]
    for h in heights:
        ycum.append(ycum[-1] + h)
    for (x0, y0, x1, y1), png in pics:
        r = max(0, max(i for i, y in enumerate(ycum[:-1]) if y <= y0) if any(y <= y0 for y in ycum[:-1]) else 0)
        c = max(i for i, b in enumerate(bounds[:-1]) if b <= x0 + 0.01) if bounds[0] <= x0 else 0
        sheet.pictures.append((r, y0 - ycum[r], c, x0 - bounds[c], x1 - x0, y1 - y0, png))
    return sheet


# --------------------------------------------------------------------------- output
def _col_width_units(pt: float) -> float:
    px = pt * 96 / 72
    return round(px / 12, 2) if px < 12 else round((px - 5) / 7, 2)


def build_exact_sheets(pages: List[PageResult], single: bool) -> List[ExactSheet]:
    """One sheet per page, or every page stacked on one sheet."""
    pages = [p for p in pages if p.kind != "error" or p.segments]
    if single:
        return [layout_pages(pages, "All pages")]
    return [layout_pages([p], f"Page {p.number}") for p in pages]


def write_exact_xlsx(sheets: List[ExactSheet], convert_types: bool, date_order: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(sh.name[:31])
        ws.sheet_view.showGridLines = False
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        for c, w in enumerate(sh.col_widths_pt, start=1):
            ws.column_dimensions[get_column_letter(c)].width = _col_width_units(w)
        for r, h in enumerate(sh.row_heights_pt, start=1):
            ws.row_dimensions[r].height = round(h, 2)
        merged: set = set()
        for r0, c0, r1, c1, color in sh.fills:
            fill = PatternFill("solid", start_color=color)
            for r in range(r0 + 1, r1 + 2):
                for c in range(c0 + 1, c1 + 2):
                    ws.cell(row=r, column=c).fill = fill
        for r, c0, c1, color in sh.rules:
            side = Side(style="thin", color=color)
            for c in range(c0 + 1, c1 + 2):
                cell = ws.cell(row=r + 1, column=c)
                b = cell.border
                cell.border = Border(left=b.left, right=b.right, top=b.top, bottom=side)
        for p in sh.cells:
            s = p.seg
            if convert_types:
                value, fmt = display_value(s.text, date_order, ocr=sh.ocr)
            else:
                value, fmt = s.text, None
            row, col = p.row + 1, p.col + 1          # openpyxl rows/columns are 1-based
            row_end, col_end = p.row_end + 1, p.col_end + 1
            cell = ws.cell(row=row, column=col, value=value if value != "" else None)
            if fmt:
                cell.number_format = fmt
            low = sh.ocr and s.text.strip() and s.conf < LOW_CONF
            cell.font = Font(name=FONT_NAME, size=s.size, bold=s.bold,
                             color=s.color if s.color else None)
            fill = CHECK_FILL if low else s.fill
            cell.alignment = Alignment(horizontal=p.align, vertical="top",
                                       wrap_text=bool(s.boxed and (col_end > col or row_end > row
                                                                   or len(s.text) > 25)))
            if low:
                cell.comment = Comment(f"Read from a scan with low confidence ({s.conf:.0f}%). "
                                       "Check it against the PDF.", "PDF to Excel")
            span = [(r, c) for r in range(row, row_end + 1) for c in range(col, col_end + 1)]
            if (row_end > row or col_end > col) and not (set(span) & merged):
                ws.merge_cells(start_row=row, start_column=col, end_row=row_end, end_column=col_end)
                merged.update(span)
            for (r, c) in span:
                target = ws.cell(row=r, column=c)
                if fill:
                    target.fill = PatternFill("solid", start_color=fill)
                if s.boxed:
                    target.border = Border(
                        left=THIN if c == col else None, right=THIN if c == col_end else None,
                        top=THIN if r == row else None, bottom=THIN if r == row_end else None)
        for r in sh.page_break_rows:
            ws.row_breaks.append(Break(id=r))            # break after 1-based row r (= index r-1)
        for r, dy, c, dx, w, h, png in sh.pictures:
            img = XLImage(io.BytesIO(png))
            emu = 12700                                    # EMU per point
            marker = AnchorMarker(col=c, colOff=int(dx * emu), row=r, rowOff=int(dy * emu))
            img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(int(w * emu), int(h * emu)))
            ws.add_image(img)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def grid_rows(sheet: ExactSheet) -> List[List[str]]:
    """Plain text grid (for CSV, JSON and the preview)."""
    n_rows = len(sheet.row_heights_pt)
    n_cols = len(sheet.col_widths_pt)
    grid = [[""] * n_cols for _ in range(n_rows)]
    for p in sheet.cells:
        if p.row < n_rows and p.col < n_cols:
            grid[p.row][p.col] = p.seg.text
    return grid


def exact_preview(sheets: List[ExactSheet], max_rows: int = 80) -> list:
    out = []
    for sh in sheets:
        grid = grid_rows(sh)
        flags = {}
        nums = {}
        for p in sh.cells:
            if sh.ocr and p.seg.text.strip() and p.seg.conf < LOW_CONF:
                flags[(p.row, p.col)] = 1
            elif p.seg.bold:
                flags[(p.row, p.col)] = 2
            nums[(p.row, p.col)] = p.align == "right"
        # drop rows/columns that are empty everywhere so the preview stays readable
        rows = [i for i, r in enumerate(grid) if any(v.strip() for v in r)]
        cols = [j for j in range(len(grid[0]) if grid else 0) if any(grid[i][j].strip() for i in rows)]
        out.append({
            "name": sh.name,
            "rows": [[{"v": grid[i][j], "f": flags.get((i, j), 0), "n": nums.get((i, j), False)}
                      for j in cols] for i in rows[:max_rows]],
            "total_rows": len(rows),
        })
    return out
